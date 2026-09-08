#!/usr/bin/env python3
"""Prüft die Schwärzung in s2_aus_vorlage.py: greift sie, und schlägt sie laut
fehl statt still, wenn die Vorlage sich verschiebt?

Das ist der Test zur einen Stelle, an der ein unbemerkter Fehler direkt in ein
öffentliches Repo schreibt. Er nennt selbst keine behördenspezifischen Werte --
er liest sie aus der lokalen Vorlage und prüft, dass keiner in der Ausgabe steht.

    python3 grist/ausgabe/test_schwaerzung.py

Braucht die Vorlage (gitignored); fehlt sie, überspringt sich der Test.
"""
import os, pathlib, shutil, subprocess, sys, tempfile

import openpyxl

HIER = pathlib.Path(__file__).resolve().parent
SKRIPT = HIER / "s2_aus_vorlage.py"
VORLAGE = HIER.parents[1] / "ReisekostenabrechnungFINAL.xlsx"

if not VORLAGE.exists():
    sys.exit(f"Übersprungen: {VORLAGE.name} liegt nicht im Repo-Wurzelverzeichnis.")

def lauf(vorlage):
    return subprocess.run([sys.executable, str(SKRIPT)], capture_output=True, text=True,
                          env={**os.environ, "S2_VORLAGE": str(vorlage)})


def regeln():
    """Die Zelladressen aus SCHWAERZEN, ohne das Skript zu importieren (es
    rendert beim Import bereits und würde bei Manipulation aussteigen)."""
    aus, sammeln = [], False
    for zeile in SKRIPT.read_text(encoding="utf-8").splitlines():
        if zeile.startswith("SCHWAERZEN = {"):
            sammeln = True
        elif sammeln and zeile.startswith("}"):
            break
        elif sammeln and zeile.strip().startswith("("):
            r, c = zeile.strip().split(")")[0].lstrip("(").split(",")
            aus.append((int(r), int(c)))
    return aus


ZELLEN = regeln()
assert ZELLEN, "keine Regeln in SCHWAERZEN gefunden"

with tempfile.TemporaryDirectory() as tmp:
    tmp = pathlib.Path(tmp)

    # 1. Normalfall: läuft durch, und kein Originalwert steht in der Ausgabe.
    e = lauf(VORLAGE)
    assert e.returncode == 0, f"Normallauf gescheitert:\n{e.stderr}"
    ws = openpyxl.load_workbook(VORLAGE)["S2"]
    originale = [str(ws.cell(r, c).value).rstrip() for r, c in ZELLEN]
    for wert in originale:
        assert wert not in e.stdout, f"Originalwert steht in der Ausgabe: {wert[:20]!r}…"
    assert e.stdout.count("_____") >= len(ZELLEN), "Platzhalter fehlen in der Ausgabe"
    print(f"  [OK] Normallauf: {len(ZELLEN)} Zellen geschwärzt, keine Originalwerte in der Ausgabe")

    # 2. Vorlage verschoben: Zelle hat anderen Inhalt -> lauter Abbruch.
    verschoben = tmp / "verschoben.xlsx"
    shutil.copy(VORLAGE, verschoben)
    wb = openpyxl.load_workbook(verschoben)
    wb["S2"].cell(*ZELLEN[0]).value = "irgendetwas anderes"
    wb.save(verschoben)
    e = lauf(verschoben)
    assert e.returncode != 0, "geänderte Zelle blieb unbemerkt — genau der stille Fehlschlag"
    assert "Fingerabdruck" in e.stderr, e.stderr
    assert not e.stdout, "trotz Abbruch wurde HTML ausgegeben"
    print("  [OK] Geänderte Zelle: Abbruch statt stiller Ausgabe")

    # 3. Originalwert steht zusätzlich in einer nicht geschwärzten Zelle.
    doppelt = tmp / "doppelt.xlsx"
    shutil.copy(VORLAGE, doppelt)
    wb = openpyxl.load_workbook(doppelt)
    verbunden = {(r, c) for m in wb["S2"].merged_cells.ranges
                 for r in range(m.min_row, m.max_row + 1)
                 for c in range(m.min_col, m.max_col + 1)}
    frei = next((r, c) for r in range(1, 49) for c in range(1, 8)
                if (r, c) not in ZELLEN and (r, c) not in verbunden
                and not wb["S2"].cell(r, c).value)
    wb["S2"].cell(*frei).value = originale[-1]
    wb.save(doppelt)
    e = lauf(doppelt)
    assert e.returncode != 0, "verdoppelter Originalwert blieb unbemerkt"
    assert not e.stdout, "trotz Abbruch wurde HTML ausgegeben"
    print("  [OK] Wert an zweiter Stelle: Abbruch statt Leak")

print("\nSchwärzung hält.")
