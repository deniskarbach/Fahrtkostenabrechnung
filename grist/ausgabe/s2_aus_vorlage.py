#!/usr/bin/env python3
"""Erzeugt den S2-Tabellenblock von ausdruck.html aus dem XLSX-Blatt "S2".

S2 ist ein amtliches Formular -- Ränder, Füllungen, Schriftgrößen,
Zeilenhöhen und Zellverbünde werden deshalb Zelle für Zelle aus der Vorlage
übernommen statt nachgebaut. Ändert sich die Vorlage, dieses Skript erneut
laufen lassen und die Ausgabe in ausdruck.html zwischen <table class="s2">
und </table> einsetzen -- nicht von Hand nachziehen.

    python3 grist/ausgabe/s2_aus_vorlage.py > /tmp/s2.html

Braucht ReisekostenabrechnungFINAL.xlsx im Repo-Wurzelverzeichnis. Die Datei
ist per .gitignore ausgeschlossen (öffentliches Repo), behördenspezifische
Zellen werden beim Erzeugen geschwärzt -- siehe SCHWAERZEN.

    python3 grist/ausgabe/s2_aus_vorlage.py --fingerabdruecke   # nach Vorlagenwechsel
"""
import hashlib
import os
import pathlib
import sys

import openpyxl

# S2_VORLAGE überschreibt den Pfad — braucht test_schwaerzung.py, um die
# Wächter gegen manipulierte Kopien der Vorlage laufen zu lassen.
VORLAGE = pathlib.Path(os.environ.get("S2_VORLAGE") or
                       pathlib.Path(__file__).resolve().parents[2]
                       / "ReisekostenabrechnungFINAL.xlsx")
WB = openpyxl.load_workbook(VORLAGE)
WS = WB["S2"]
SPALTEN = list("ABCDEFG")
DEF_W = 12.6328125
SATZ = 145.79   # gemessene Breite des Originaldrucks (PDF der Vorlage,
                # Skalierung 100%): 17,46mm .. 163,25mm ab Blattkante.
                # NICHT auf den Satzspiegel normiert -- sonst stimmen die
                # Zeilenumbrüche im Fließtext nicht mit dem Formular überein.
PT_MM = 25.4 / 72

# Behördenspezifisches bleibt Platzhalter — das Repo ist öffentlich. Adressiert
# wird über die Zelle, nicht über ihren Inhalt: so steht der zu schwärzende Wert
# nirgends in dieser Datei. Der Fingerabdruck (sha256, 12 Zeichen, über den
# rechts beschnittenen Zellwert) sichert die Adresse ab -- verschiebt sich die
# Vorlage um eine Zeile, bricht der Lauf ab, statt den echten Wert still in die
# öffentliche Ausgabe zu schreiben. Neue Vorlage: mit --fingerabdruecke die
# Werte der hier eingetragenen Zellen neu ausgeben lassen und prüfen.
SCHWAERZEN = {
    (34, 1): ("81e3e161d025", "Abteilung: __________"),
    (34, 4): ("3df650914e3b", "_______________, _______________"),
    (37, 1): ("719b8d22e6d7",
              "1. KK anweisen,                         \u20ac bei: _____________________"),
    (45, 3): ("ed6bb6080fdc", "_____________________________"),
}

GESCHWAERZT = {}   # (Zeile, Spalte) -> Originalwert, für die Schlusskontrolle


def schwaerzen(r, c, wert):
    """Ersetzt den Inhalt einer als behördenspezifisch markierten Zelle. Bricht
    ab, sobald die Zelle nicht mehr den erwarteten Inhalt hat -- der stille
    Fehlschlag ist hier das eigentliche Risiko, nicht der laute."""
    erwartet, ersatz = SCHWAERZEN[(r, c)]
    ist = hashlib.sha256(wert.encode()).hexdigest()[:12]
    if ist != erwartet:
        sys.exit(f"Zelle {WS.cell(r, c).coordinate}: Inhalt weicht von der "
                 f"hinterlegten Schwärzungsregel ab (erwartet {erwartet}, ist {ist}).\n"
                 f"Vorlage geändert? Regel in SCHWAERZEN prüfen und den "
                 f"Fingerabdruck mit --fingerabdruecke neu bestimmen. "
                 f"Abbruch, damit nichts Behördenspezifisches in die "
                 f"öffentliche Ausgabe gerät.")
    GESCHWAERZT[(r, c)] = wert
    return ersatz


if "--fingerabdruecke" in sys.argv:
    # Gibt Adresse, Fingerabdruck und Inhalt der in SCHWAERZEN eingetragenen
    # Zellen aus -- nur lokal, die Vorlage liegt ja ohnehin nur lokal.
    for (zr, zc) in sorted(SCHWAERZEN):
        w = str(WS.cell(zr, zc).value or "").rstrip()
        print(f"({zr:3d}, {zc}): {hashlib.sha256(w.encode()).hexdigest()[:12]}   {w!r}")
    sys.exit(0)

STAERKE = {"thin": "var(--rahmen)", "medium": "var(--rahmen-breit)"}

# Zellen, die das Skript im Widget füllt (Zeile, Spaltenindex) -> id
IDS = {(5, 3): "k-s2-iban", (7, 3): "k-s2-bic"}


def breite(sp):
    d = WS.column_dimensions.get(sp)
    return d.width if d and d.width else DEF_W


def hoehe(r):
    d = WS.row_dimensions.get(r)
    return d.height if d and d.height else 12.75


# ---------------------------------------------------------------- Merges
ANKER = {}      # (row, col_index) -> (rowspan, colspan)
BELEGT = set()  # von einem Merge überdeckte Zellen
for m in WS.merged_cells.ranges:
    ANKER[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1,
                                     m.max_col - m.min_col + 1)
    for r in range(m.min_row, m.max_row + 1):
        for c in range(m.min_col, m.max_col + 1):
            if (r, c) != (m.min_row, m.min_col):
                BELEGT.add((r, c))


def rahmen(r, c, rs, cs):
    """Rahmen einer (ggf. verbundenen) Zelle: je Seite aus den Randzellen des
    verbundenen Bereichs, nicht nur aus der Ankerzelle."""
    seiten = {}
    for spalte in range(c, c + cs):                      # oben / unten
        for seite, zeile in (("top", r), ("bottom", r + rs - 1)):
            stil = getattr(WS.cell(zeile, spalte).border, seite).style
            if stil:
                seiten[seite] = stil
    for zeile in range(r, r + rs):                       # links / rechts
        for seite, spalte in (("left", c), ("right", c + cs - 1)):
            stil = getattr(WS.cell(zeile, spalte).border, seite).style
            if stil:
                seiten[seite] = stil
    return seiten


# ---------------------------------------------------------------- Ausgabe
aus = []
gesamt = sum(breite(s) for s in SPALTEN)
aus.append('  <table class="s2">')
aus.append("    <colgroup>")
for s in SPALTEN:
    aus.append(f'      <col style="width:{breite(s) / gesamt * SATZ:.2f}mm">')
aus.append("    </colgroup>")

for r in range(1, 49):
    aus.append(f'    <tr style="height:{hoehe(r) * PT_MM:.2f}mm">')
    for ci, sp in enumerate(SPALTEN, start=1):
        if (r, ci) in BELEGT:
            continue
        rs, cs = ANKER.get((r, ci), (1, 1))
        zelle = WS.cell(r, ci)
        stil, attr = [], []
        if rs > 1:
            attr.append(f'rowspan="{rs}"')
        if cs > 1:
            attr.append(f'colspan="{cs}"')

        for seite, art in rahmen(r, ci, rs, cs).items():
            stil.append(f"border-{seite}:{STAERKE.get(art, 'var(--rahmen)')}")

        fuell = zelle.fill
        if fuell and fuell.patternType and fuell.fgColor and fuell.fgColor.rgb:
            stil.append(f"background:#{str(fuell.fgColor.rgb)[2:]}")

        wert = zelle.value
        # rstrip: die Vorlage hat in mehreren Zellen lange Leerzeichenketten am
        # Ende. Als &nbsp; ausgegeben schieben sie den Text weit über den
        # Seitenrand -- der Browser skaliert dann das GANZE Dokument klein.
        wert = "" if wert is None else str(wert).rstrip()
        if wert.startswith("="):        # Formeln der Vorlage: leer lassen
            wert = ""
        if (r, ci) in SCHWAERZEN:
            wert = schwaerzen(r, ci, wert)

        if wert:
            f, a = zelle.font, zelle.alignment
            stil.append(f"font-size:{f.size:g}pt")
            if f.bold:
                stil.append("font-weight:bold")
            if a.horizontal in ("center", "right", "left"):
                stil.append(f"text-align:{a.horizontal}")
            if a.vertical in ("top", "center", "bottom"):
                stil.append("vertical-align:" +
                            {"center": "middle"}.get(a.vertical, a.vertical))
            if a.wrapText:
                stil.append("white-space:normal")
            else:
                # Excel laesst Text in leere Nachbarzellen ueberlaufen und kappt
                # ihn erst an der naechsten belegten Zelle -- genau das hier.
                nachbar = ci + cs
                if nachbar <= len(SPALTEN) and WS.cell(r, nachbar).value not in (None, ""):
                    stil.append("overflow:hidden")

        inhalt = (wert.replace("&", "&amp;").replace("<", "&lt;")
                      .replace(">", "&gt;").replace("  ", "&nbsp;&nbsp;"))
        if (r, ci) in IDS:
            attr.insert(0, f'id="{IDS[(r, ci)]}"')
        offen = "<td" + ("".join(" " + x for x in attr))
        if stil:
            offen += ' style="' + ";".join(stil) + '"'
        aus.append(f"      {offen}>{inhalt}</td>")
    aus.append("    </tr>")
aus.append("  </table>")

# ------------------------------------------------- Kontrolle vor der Ausgabe
# Zwei Dinge müssen stimmen, sonst wird nichts gedruckt: jede Regel muss eine
# Zelle erwischt haben (eine von einem Merge überdeckte Zelle besucht die
# Schleife nie), und kein geschwärzter Originalwert darf anderswo im Blatt
# stehen geblieben sein.
fehlend = sorted(SCHWAERZEN.keys() - GESCHWAERZT.keys())
if fehlend:
    sys.exit(f"Schwärzungsregeln liefen ins Leere für Zellen {fehlend} — "
             f"Zeilenbereich oder Zellverbünde der Vorlage geändert? Abbruch.")

fertig = "\n".join(aus)
uebrig = sorted({WS.cell(r, c).coordinate for (r, c), wert in GESCHWAERZT.items()
                 if wert in fertig})
if uebrig:
    sys.exit(f"Ein geschwärzter Wert steht noch an anderer Stelle im Blatt "
             f"(Ursprung: {uebrig}). Betroffene Zelle in SCHWAERZEN ergänzen. Abbruch.")

print(fertig)
